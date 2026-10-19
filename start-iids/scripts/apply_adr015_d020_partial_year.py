"""One-off data migration for ADR-015 — D020 partial production year 2025.

D020 (Viano) ceased production in July 2025 (reconversion to a logistics hub).
Its 2025 is therefore a PARTIAL period (Jan -> cessation), not a full year. Per
ADR-015 (option A, homologous reference), D020/2025's module baseline becomes the
SAME Jan->cessation window of 2022, not the annual 2022 row. Because reference and
current period use the same monthly window, the intensive ratios (Phi, Psi,
TSI_abs, TSI_rel) are invariant, while the extensive quantities (f_i, SA, Ex_ref,
production) scale to the period and are flagged NON_COMPARABLE with full years.

The provisional period values here are generated ONLY to exercise the model
(project owner, 2026-10-19); they are to be replaced by the real measured period
data collected via `RP7.3_D020_partial_year_data_collection.xlsx`. They are NOT a
flat month-fraction of the annual values (that would be forbidden reproportioning)
but the Jan->Jul share of a seasonal monthly production profile.

Writes both mirror copies (repo root + start-iids/data/reference) of the data
collection workbook and the calculation log (bumped to beta-1.2). Idempotent:
reads the archived beta-1.1 fixtures as the deterministic source.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

from src.engines.ecofa.formulas import compute_f_econ
from src.engines.efa.formulas import compute_f_env
from src.engines.sfa.formulas import compute_f_soc
from src.engines.tei.formulas import compute_f_tech

ROOT = Path(__file__).resolve().parents[1]              # start-iids/
REPO_ROOT = ROOT.parent                                 # Progetto-MIMIT-START/
DATA = ROOT / "data" / "reference"
FIX = ROOT / "tests" / "fixtures" / "rp73"

WB_NAME = "RP7.3_data_collection_20232025.xlsx"
LOG_NAME = "RP7.3_calculation_log.xlsx"
WB_COPIES = [DATA / WB_NAME, REPO_ROOT / WB_NAME]
LOG_COPIES = [DATA / LOG_NAME, REPO_ROOT / LOG_NAME]

# --- ADR-015 parameters ---
# Seasonal monthly production profile (ceramic; low August plant shutdown), sums to 1.
MONTHLY = [.090, .090, .095, .092, .092, .088, .080, .030, .090, .095, .090, .068]
CESSATION_MONTH = 7                                     # July 2025
FRAC = round(sum(MONTHLY[:CESSATION_MONTH]), 4)         # Jan-Jul seasonal share = 0.627 (NOT 7/12)
MONTHS_ACTIVE = CESSATION_MONTH
PERIOD_2025 = ("2025-01-01", "2025-07-31")
PERIOD_2022H = ("2022-01-01", "2022-07-31")
TODAY = "2026-10-19"
NEW_VERSION = "beta-1.2"

PLANT_PROD = {"D020": 3_800_000, "D060": 6_400_000, "D240": 5_100_000}
W = {"env": 0.3661, "econ": 0.1451, "soc": 0.0955, "tech": 0.3934}
EL_EX, GAS_EX = 3.6, 42

TEI_KEYS = ["loss_MTS_MJ", "loss_MTO_MJ", "inv_MJ", "qual_MTS_MJ", "qual_MTO_MJ"]
EFA_KEYS = ["RI_MJ", "IEQ_MJ", "WEX_MJ", "CIRC_MJ"]
ECO_KEYS = ["VA_MJ", "econ_in_MJ", "INV_MJ"]
SFA_KEYS = ["SV_MJ", "train_MJ", "lost_MJ", "CO2_MJ"]


def _sheet_map(wb, sheet):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    return hdr, {(r[0], r[1]): dict(zip(hdr, r)) for r in rows[1:]}


def build_workbook():
    src = FIX / "RP7.3_data_collection_beta-1.1.xlsx"
    wb = openpyxl.load_workbook(src)
    _, energy = _sheet_map(wb, "Energia")
    _, tei = _sheet_map(wb, "Moduli_TEI")
    _, efa = _sheet_map(wb, "Moduli_EFA")
    _, eco = _sheet_map(wb, "Moduli_EcoFA")
    _, sfa = _sheet_map(wb, "Moduli_SFA")

    # 1) Unita -> per-year production (plant, year, produzione_m2, descrizione)
    desc = {r[0]: r[1] for r in list(wb["Unita"].iter_rows(values_only=True))[1:]}
    wb.remove(wb["Unita"])
    ws = wb.create_sheet("Unita"); _move_first(wb, ws)
    ws.append(["plant", "year", "produzione_m2", "descrizione"])
    for plant in PLANT_PROD:
        for year in (2022, 2023, 2024, 2025):
            prod = PLANT_PROD[plant]
            if plant == "D020" and year == 2025:
                prod = round(PLANT_PROD[plant] * FRAC)        # period production
            ws.append([plant, year, prod, desc.get(plant, "")])

    # 2) Periodo sheet (plant, year, period_start, period_end, months_active)
    per = wb.create_sheet("Periodo")
    per.append(["plant", "year", "period_start", "period_end", "months_active"])
    for plant in PLANT_PROD:
        for year in (2022, 2023, 2024, 2025):
            if plant == "D020" and year == 2025:
                per.append([plant, year, PERIOD_2025[0], PERIOD_2025[1], MONTHS_ACTIVE])
            else:
                per.append([plant, year, f"{year}-01-01", f"{year}-12-31", 12])

    # 3) Riferimento sheet (homologous baseline for D020/2025): long format
    rif = wb.create_sheet("Riferimento")
    rif.append(["plant", "year", "ref_plant", "ref_year", "ref_period_start",
                "ref_period_end", "module", "term", "value_MJ"])
    homol = {}
    for mod, keys, base in [("TEI-J", TEI_KEYS, tei[("D020", 2022)]),
                            ("EFA-J", EFA_KEYS, efa[("D020", 2022)]),
                            ("EcoFA-J", ECO_KEYS, eco[("D020", 2022)]),
                            ("SFA-J", SFA_KEYS, sfa[("D020", 2022)])]:
        for term in keys:
            val = round(base[term] * FRAC)
            homol[(mod, term)] = val
            rif.append(["D020", 2025, "D020", 2022, PERIOD_2022H[0], PERIOD_2022H[1],
                        mod, term, val])

    # 4) D020/2025 Energia + Moduli -> period values (FRAC * 2025 annual)
    def _update_row(sheet, keys):
        ws = wb[sheet]
        base = {"Moduli_TEI": tei, "Moduli_EFA": efa, "Moduli_EcoFA": eco,
                "Moduli_SFA": sfa, "Energia": energy}[sheet][("D020", 2025)]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "D020" and row[1].value == 2025:
                hdr = [c.value for c in ws[1]]
                for term in keys:
                    row[hdr.index(term)].value = round(base[term] * FRAC)
    _update_row("Energia", ["V_gas_Nm3", "E_el_kWh"])
    _update_row("Moduli_TEI", TEI_KEYS)
    _update_row("Moduli_EFA", EFA_KEYS)
    _update_row("Moduli_EcoFA", ECO_KEYS)
    _update_row("Moduli_SFA", SFA_KEYS)

    wb.save(WB_COPIES[0])                        # save once, then byte-copy -> identical md5
    for dst in WB_COPIES[1:]:
        shutil.copyfile(WB_COPIES[0], dst)
    print(f"workbook beta-1.2 written to {len(WB_COPIES)} copies (FRAC={FRAC}, months={MONTHS_ACTIVE})")
    return energy, tei, efa, eco, sfa, homol


def _move_first(wb, ws):
    wb._sheets.remove(ws); wb._sheets.insert(0, ws)


def compute_d020_2025_period(energy, tei, efa, eco, sfa, homol):
    """Compute the D020/2025 period outputs for the log (option A, homologous ref)."""
    def cur(d, keys):
        return {k: d[("D020", 2025)][k] * FRAC for k in keys}
    e = cur(energy, ["V_gas_Nm3", "E_el_kWh"]); t = cur(tei, TEI_KEYS)
    f = cur(efa, EFA_KEYS); c = cur(eco, ECO_KEYS); s = cur(sfa, SFA_KEYS)
    hb = lambda mod, term: homol[(mod, term)]
    f_env = compute_f_env(f["RI_MJ"], hb("EFA-J", "RI_MJ"), f["CIRC_MJ"], hb("EFA-J", "CIRC_MJ"),
                          f["IEQ_MJ"], hb("EFA-J", "IEQ_MJ"), f["WEX_MJ"], hb("EFA-J", "WEX_MJ")) / 1000
    f_econ = compute_f_econ(c["VA_MJ"], hb("EcoFA-J", "VA_MJ"), c["econ_in_MJ"], hb("EcoFA-J", "econ_in_MJ"),
                            c["INV_MJ"], hb("EcoFA-J", "INV_MJ")) / 1000
    f_soc = compute_f_soc(s["SV_MJ"], hb("SFA-J", "SV_MJ"), s["train_MJ"], hb("SFA-J", "train_MJ"),
                          s["lost_MJ"], hb("SFA-J", "lost_MJ"), s["CO2_MJ"], hb("SFA-J", "CO2_MJ")) / 1000
    f_tech = compute_f_tech(ex_loss_base_mts_mj=hb("TEI-J", "loss_MTS_MJ"),
                            ex_loss_base_mto_mj=hb("TEI-J", "loss_MTO_MJ"),
                            ex_loss_mts_mj=t["loss_MTS_MJ"], ex_loss_mto_mj=t["loss_MTO_MJ"],
                            ex_inv_mj=t["inv_MJ"], ex_qual_mts_mj=t["qual_MTS_MJ"],
                            ex_qual_mto_mj=t["qual_MTO_MJ"]) / 1000
    sa_raw = f_env + f_econ + f_soc + f_tech
    sa_w = W["env"] * f_env + W["econ"] * f_econ + W["soc"] * f_soc + W["tech"] * f_tech
    ex_ref = (e["E_el_kWh"] * EL_EX + e["V_gas_Nm3"] * GAS_EX) / 1000
    phi = sa_w / ex_ref
    psi = 0.164
    tsi_abs = 0.5 * phi + 0.5 * psi
    return {"f_env": f_env, "f_econ": f_econ, "f_soc": f_soc, "f_tech": f_tech,
            "SA_raw": sa_raw, "SA_w": sa_w, "Ex_ref": ex_ref, "Phi": phi,
            "Psi": psi, "TSI_abs": tsi_abs}


def build_log(d020_2025):
    src = FIX / "RP7.3_calculation_log_beta-1.1.xlsx"
    wb = openpyxl.load_workbook(src)
    ws = wb["calculation_log"]
    hdr = [c.value for c in ws[1]]
    i_var, i_out, i_formula, i_date, i_ver = (hdr.index("variable"), hdr.index("output"),
                                              hdr.index("formula"), hdr.index("date"),
                                              hdr.index("script_version"))
    tsi_abs_2023 = {}
    for row in ws.iter_rows(min_row=2):
        if row[2].value == "D020" and row[3].value == 2023 and row[i_var].value == "TSI_abs":
            tsi_abs_2023["D020"] = row[i_out].value
    for row in ws.iter_rows(min_row=2):
        row[i_ver].value = NEW_VERSION                        # bump every row to beta-1.2
        if row[2].value == "D020" and row[3].value == 2025:
            var = row[i_var].value
            row[i_date].value = TODAY
            if var in d020_2025:
                row[i_out].value = round(d020_2025[var], 4)
            elif var == "TSI_rel":
                row[i_out].value = round(d020_2025["TSI_abs"] / tsi_abs_2023["D020"], 4)
            # annotate homologous reference in the formula strings
            if var == "f_env":
                row[i_formula].value = "(RI_ref[2022 Jan-Jul]-RI)+(CIRC-CIRC_ref)-(IEQ-IEQ_ref)-(WEX-WEX_ref) [ADR-015 option A]"
            elif var in ("f_econ", "f_soc", "f_tech"):
                row[i_formula].value = (row[i_formula].value or "") + " [ADR-015 option A: ref=2022 Jan-Jul]"
            elif var == "Ex_ref":
                row[i_formula].value = "Ex_el+Ex_fuel (period Jan-Jul 2025)"
    wb.save(LOG_COPIES[0])                        # save once, then byte-copy -> identical md5
    for dst in LOG_COPIES[1:]:
        shutil.copyfile(LOG_COPIES[0], dst)
    print(f"calculation log {NEW_VERSION} written to {len(LOG_COPIES)} copies")


def main():
    energy, tei, efa, eco, sfa, homol = build_workbook()
    d020 = compute_d020_2025_period(energy, tei, efa, eco, sfa, homol)
    print("D020/2025 period:", {k: round(v, 4) for k, v in d020.items()},
          "TSI_rel=", round(d020["TSI_abs"] / 0.0867, 4))
    build_log(d020)


if __name__ == "__main__":
    main()
